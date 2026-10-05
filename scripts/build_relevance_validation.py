#!/usr/bin/env python3
"""Arma una muestra estratificada para validar a mano el filtro de pertinencia.

Toma todos los registros reunidos (maestro, revisión y rechazados), los
clasifica con el filtro actual y sortea registros de cuatro estratos: fuente
(OpenAlex / CONICET) × decisión del filtro (incluido / no incluido). La planilla
no muestra la decisión del filtro, para no condicionar la lectura.

Salida: data/validacion/muestra_pertinencia.xlsx. Completar la columna
«pertinente» (si / no / dudoso) y correr scripts/evaluate_relevance_validation.py.

Criterio sugerido: «si» cuando el objeto de estudio es la dirección, gestión,
gobierno o liderazgo de escuelas (niveles inicial, primario o secundario), sus
actores directivos o las políticas que los regulan.
"""
from __future__ import annotations

import argparse
import random
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

import relevance_filter as rf  # noqa: E402

OUT = ROOT / "data" / "validacion" / "muestra_pertinencia.xlsx"
ALLOCATION = {("OpenAlex", "incluido"): 60, ("OpenAlex", "no incluido"): 40,
              ("CONICET", "incluido"): 50, ("CONICET", "no incluido"): 50}
FIELDS = ["id_muestra", "record_id", "fuente", "anio", "titulo", "resumen", "palabras_clave",
          "tipo", "url", "pertinente", "marcado_por", "nota"]

# Títulos que nombran explícitamente la dirección, gestión o liderazgo escolar:
# se marcan «si» automáticamente (editable) para revisar a mano solo el resto.
OBVIOUS_TITLE = re.compile(r"""
 \b(gestion|direccion|liderazgo|conduccion|administracion|gobierno|gestao|direcao|lideranca)\s+
   (escolar\w*|educativ\w*|educacional|directiv\w*|pedagogic\w*|democratica|de\s+(la\s+)?escuela\w*|da\s+escola)
|\b(director\w*|directiv\w*|diretor\w*|gestor\w*|supervisor\w*)\b
|\bequipos?\s+directivos?\b
|\b(school|educational|instructional|pedagogical|distributed|principal)\s+(leader\w*|management|administration|governance)
|\bprincipals?\b.{0,40}\b(school|teacher|leadership)|\bheadteacher\w*|\bprincipalship\b
|\bschool\s+(heads?|principals?)\b|\bschool[- ]based\s+management\b
""", re.X)
HIGHER_ED_TITLE = re.compile(r"\b(universi\w*|higher\s+education|educacion\s+superior|posgrado|postgrado|"
                             r"academic\s+leadership|department\s+heads?|faculty|surgeons?|hospital|medic\w*)\b")


def obvious_by_title(title: str) -> bool:
    if rf.LIBRARY_GUIDE_RE.match(str(title or "")):
        return False
    t = rf.strip_accents(str(title or "").lower())
    return bool(OBVIOUS_TITLE.search(t)) and not HIGHER_ED_TITLE.search(t)


def prefill(path: Path) -> int:
    """Marca «si» los títulos obvios sin tocar las marcas ya cargadas."""
    wb = openpyxl.load_workbook(path)
    ws = wb["Muestra"]
    header = [c.value for c in ws[1]]
    if "marcado_por" not in header:
        ws.insert_cols(header.index("pertinente") + 2)
        ws.cell(row=1, column=header.index("pertinente") + 2, value="marcado_por")
        header = [c.value for c in ws[1]]
    col_t, col_p, col_m = (header.index(x) + 1 for x in ("titulo", "pertinente", "marcado_por"))
    count = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, col_p).value in (None, "") and obvious_by_title(ws.cell(r, col_t).value):
            ws.cell(r, col_p, "si")
            ws.cell(r, col_m, "automático: título explícito")
            count += 1
    ws.auto_filter.ref = ws.dimensions  # filtrar «pertinente» vacío para ver lo pendiente
    wb.save(path)
    return count


def final_status(row: dict) -> str:
    status = rf.classify_relevance(row)[0]
    if status == "revisar":
        status = {"promover": "alta", "descartar": "rechazada"}.get(rf.second_review(row)[0], "revisar")
    return "incluido" if status == "alta" else "no incluido"


def candidates() -> list[dict]:
    rows: list[dict] = []
    for name in ("master_records", "review_records", "rejected_records"):
        _, part = rf.read_csv_if_exists(ROOT / "data" / f"{name}.csv")
        rf.append_unique(rows, part)
    return rows


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--seed", type=int, default=20261005)
    parser.add_argument("--out", type=Path, default=OUT)
    parser.add_argument("--solo-premarcar", action="store_true",
                        help="no sortear: solo marcar títulos obvios en la planilla existente")
    args = parser.parse_args(argv)
    if args.solo_premarcar:
        print(f"{prefill(args.out)} registros marcados por título en {args.out}")
        return

    strata: dict[tuple[str, str], list[dict]] = {key: [] for key in ALLOCATION}
    for row in candidates():
        source = "CONICET" if "conicet" in rf.source_name(row).lower() else "OpenAlex"
        strata[(source, final_status(row))].append(row)

    rng = random.Random(args.seed)
    sample = []
    for key, size in ALLOCATION.items():
        pool = sorted(strata[key], key=lambda r: rf.record_key(r))
        for row in rng.sample(pool, min(size, len(pool))):
            sample.append((key, row))
    rng.shuffle(sample)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Muestra"
    ws.append(FIELDS)
    for i, (_, row) in enumerate(sample, 1):
        ws.append([i, rf.record_key(row), rf.source_name(row), row.get("publication_year", ""),
                   row.get("title", ""), rf.clean_abstract(row)[:1500], row.get("keywords", ""),
                   row.get("document_type", ""), row.get("url", ""), "", "", ""])
    widths = {"titulo": 60, "resumen": 90, "palabras_clave": 30, "url": 30, "nota": 30, "record_id": 22}
    for col, field in enumerate(FIELDS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = widths.get(field, 12)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    col = openpyxl.utils.get_column_letter(FIELDS.index("pertinente") + 1)
    dv = DataValidation(type="list", formula1='"si,no,dudoso"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{len(sample) + 1}")
    ws.freeze_panes = "F2"

    meta = wb.create_sheet("Estratos")
    meta.append(["record_id", "fuente", "decision_filtro"])
    for key, row in sample:
        meta.append([rf.record_key(row), *key])
    meta.sheet_state = "hidden"
    pop = wb.create_sheet("Poblacion")
    pop.append(["fuente", "decision_filtro", "registros", "muestra"])
    for key, rows in strata.items():
        pop.append([*key, len(rows), min(ALLOCATION[key], len(rows))])
    pop.sheet_state = "hidden"

    guide = wb.create_sheet("Instrucciones", 0)
    for line in [
        "Validación del filtro de pertinencia",
        "",
        "Completar en la hoja «Muestra» la columna «pertinente» con si / no / dudoso.",
        "Los títulos que nombran explícitamente la dirección o gestión escolar ya vienen marcados «si»",
        "(columna «marcado_por»: automático). Se pueden corregir si alguno no corresponde.",
        "Para ver solo lo pendiente: filtro del encabezado «pertinente» → (Vacías).",
        "Criterio: «si» cuando el objeto de estudio es la dirección, gestión, gobierno o liderazgo",
        "de escuelas (inicial, primaria, secundaria), sus actores directivos o las políticas que los regulan.",
        "La planilla no muestra qué decidió el filtro, para no condicionar la lectura.",
        "Al terminar: python scripts/evaluate_relevance_validation.py",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    guide["A1"].font = Font(bold=True, size=13)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    filled = prefill(args.out)
    print(f"Muestra de {len(sample)} registros en {args.out} ({filled} marcados por título)")
    for key, rows in strata.items():
        print(f"  {key[0]:8} {key[1]:12} población {len(rows):5}  muestra {min(ALLOCATION[key], len(rows))}")


if __name__ == "__main__":
    main()
